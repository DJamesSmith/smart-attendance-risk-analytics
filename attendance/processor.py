import csv
import logging
import re
import sys
import time
import zipfile
from datetime import date, datetime, timedelta
from datetime import time as clock_time
from functools import wraps
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

logger = logging.getLogger("attendance.processor")

REQUIRED_COLUMNS = (
    "employee_code", "employee_name", "department", "punch_date", "punch_time",
    "punch_type", "shift_start", "shift_end",
)
OPTIONAL_COLUMNS = ("device_location", "leave_type", "manager_code", "remarks")
ALLOWED_EXTENSIONS = (".csv", ".xlsx")
LEAVE_TYPES = {"CASUAL", "SICK", "EARNED", "UNPAID", "OTHER"}
GRACE_MINUTES = 10

EMPLOYEE_CODE_PATTERN = re.compile(r"^[A-Z]{3}\d{3,5}$")
ISO_DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DMY_DATE_PATTERN = re.compile(r"^\d{2}/\d{2}/\d{4}$")
TIME_PATTERN = re.compile(r"^([01]\d|2[0-3]):[0-5]\d(:[0-5]\d)?$")

DEPARTMENT_ALIASES = {
    "HR": "Human Resources",
    "IT": "IT Support",
    "IT SUPPORT": "IT Support",
    "CS": "Customer Support",
    "OPS": "Operations",
}

DEPARTMENT_HIERARCHY = {
    "Company": {
        "Operations": {"Customer Support": {}, "Field Service": {}},
        "Corporate": {"Human Resources": {}, "Finance": {}},
        "Technology": {"Engineering": {}, "IT Support": {}},
    }
}

SUMMARY_COLUMNS = (
    "employee_code", "employee_name", "department", "department_path", "punch_date",
    "first_in", "last_out", "working_hours", "late_minutes", "early_exit_minutes", "status",
)
ANOMALY_COLUMNS = ("row_number", "employee_code", "punch_date", "anomaly_type", "details")
CORRECTION_COLUMNS = (
    "employee_code", "employee_name", "department", "punch_date",
    "missing_punch", "recorded_punch", "suggested_action",)


class AttendanceFileError(Exception):
    pass


class UnsupportedFileError(AttendanceFileError):
    pass


class EmptyFileError(AttendanceFileError):
    pass


class MissingColumnsError(AttendanceFileError):
    def __init__(self, missing_columns):
        self.missing_columns = missing_columns
        super().__init__(f"Missing required columns: {', '.join(missing_columns)}")


class InvalidRowError(Exception):
    def __init__(self, row_number, field, message):
        self.row_number = row_number
        self.field = field
        super().__init__(f"{field}: {message}")


def log_execution(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        started = time.perf_counter()
        logger.info("Started %s", func.__name__)
        result = func(*args, **kwargs)
        logger.info("Finished %s in %.3f seconds", func.__name__, time.perf_counter() - started)
        return result
    return wrapper


def make_shift_rule(shift_start, shift_end, grace_minutes=GRACE_MINUTES):
    return frozenset({
        ("shift_start", shift_start),
        ("shift_end", shift_end),
        ("grace_minutes", grace_minutes),
    })


DEFAULT_SHIFT = make_shift_rule(clock_time(9, 0), clock_time(18, 0))


def find_department_path(tree, department, parents=()):
    for name, children in tree.items():
        path = parents + (name,)
        if name == department:
            return " > ".join(path)
        found = find_department_path(children, department, path)
        if found:
            return found
    return ""


def clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, clock_time):
        return value.strftime("%H:%M:%S")
    return str(value).strip()


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as file:
        for row in csv.reader(file):
            yield [clean_cell(cell) for cell in row]


def read_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for row in workbook.active.iter_rows(values_only=True):
            yield [clean_cell(cell) for cell in row]
    finally:
        workbook.close()


def stream_records(path):
    extension = Path(path).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type '{extension}'. Upload a .csv or .xlsx file.")

    rows = read_csv(path) if extension == ".csv" else read_excel(path)
    numbered_rows = ((number, row) for number, row in enumerate(rows, start=1) if any(row))

    first = next(numbered_rows, None)
    if first is None:
        raise EmptyFileError("The uploaded file is empty.")

    columns = [cell.strip().lower() for cell in first[1]]
    missing = [column for column in REQUIRED_COLUMNS if column not in columns]
    if missing:
        raise MissingColumnsError(missing)

    for row_number, row in numbered_rows:
        padded = row + [""] * (len(columns) - len(row))
        yield row_number, dict(zip(columns, padded))


def normalize_spaces(value):
    return " ".join(value.split())


def normalize_department(value):
    name = normalize_spaces(value.replace("_", " ").replace("-", " "))
    return DEPARTMENT_ALIASES.get(name.upper(), name.title())


def normalize_location(value):
    return normalize_spaces(value.replace("_", " ").replace("-", " ")).title()


def normalize_punch_type(value):
    letters = re.sub(r"[^A-Z]", "", value.upper())
    if letters[-2:] == "IN":
        return "IN"
    if letters[-3:] == "OUT":
        return "OUT"
    return ""


def parse_date(value, row_number, field):
    if ISO_DATE_PATTERN.match(value):
        date_format = "%Y-%m-%d"
    elif DMY_DATE_PATTERN.match(value):
        date_format = "%d/%m/%Y"
    else:
        raise InvalidRowError(row_number, field, f"'{value}' is not in YYYY-MM-DD or DD/MM/YYYY format")
    try:
        return datetime.strptime(value, date_format).date()
    except ValueError:
        raise InvalidRowError(row_number, field, f"'{value}' is not a real calendar date")


def parse_time(value, row_number, field):
    if not TIME_PATTERN.match(value):
        raise InvalidRowError(row_number, field, f"'{value}' is not in HH:MM or HH:MM:SS format")
    time_format = "%H:%M:%S" if len(value) == 8 else "%H:%M"
    return datetime.strptime(value, time_format).time()


def validate_row(row_number, row):
    employee_code = row.get("employee_code", "").strip().upper()
    if not EMPLOYEE_CODE_PATTERN.match(employee_code):
        raise InvalidRowError(row_number, "employee_code", f"'{employee_code}' must look like EMP1001")

    employee_name = normalize_spaces(row.get("employee_name", "")).title()
    if not employee_name:
        raise InvalidRowError(row_number, "employee_name", "is required")

    department = normalize_department(row.get("department", ""))
    if not department:
        raise InvalidRowError(row_number, "department", "is required")

    manager_code = row.get("manager_code", "").strip().upper()
    if manager_code and not EMPLOYEE_CODE_PATTERN.match(manager_code):
        raise InvalidRowError(row_number, "manager_code", f"'{manager_code}' must look like MGR101")

    leave_type = row.get("leave_type", "").strip().upper()
    if leave_type and leave_type not in LEAVE_TYPES:
        raise InvalidRowError(row_number, "leave_type", f"'{leave_type}' is not a known leave type")

    punch_date = parse_date(row.get("punch_date", "").strip(), row_number, "punch_date")
    shift_start = parse_time(row.get("shift_start", "").strip(), row_number, "shift_start")
    shift_end = parse_time(row.get("shift_end", "").strip(), row_number, "shift_end")
    if shift_end <= shift_start:
        raise InvalidRowError(row_number, "shift_end", "must be later than shift_start")

    raw_time = row.get("punch_time", "").strip()
    punch_time = None
    punch_type = ""
    if raw_time or not leave_type:
        punch_time = parse_time(raw_time, row_number, "punch_time")
        punch_type = normalize_punch_type(row.get("punch_type", ""))
        if not punch_type:
            raise InvalidRowError(row_number, "punch_type", f"'{row.get('punch_type', '')}' must be IN or OUT")

    return {
        "row_number": row_number,
        "employee_code": employee_code,
        "employee_name": employee_name,
        "department": department,
        "punch_date": punch_date,
        "punch_time": punch_time,
        "punch_type": punch_type,
        "device_location": normalize_location(row.get("device_location", "")),
        "leave_type": leave_type,
        "manager_code": manager_code,
        "remarks": normalize_spaces(row.get("remarks", "")),
        "shift": make_shift_rule(shift_start, shift_end),
    }


def minutes_between(start, end):
    difference = datetime.combine(date.min, end) - datetime.combine(date.min, start)
    return int(difference / timedelta(minutes=1))


def summarize_day(day_records):
    first = day_records[0]
    rule = dict(first["shift"])
    in_times = sorted(record["punch_time"] for record in day_records if record["punch_type"] == "IN")
    out_times = sorted(record["punch_time"] for record in day_records if record["punch_type"] == "OUT")
    leave_type = next((record["leave_type"] for record in day_records if record["leave_type"]), "")

    first_in = in_times[0] if in_times else None
    last_out = out_times[-1] if out_times else None

    delay = minutes_between(rule["shift_start"], first_in) if first_in else 0
    late_minutes = delay if delay > rule["grace_minutes"] else 0
    early_exit = minutes_between(last_out, rule["shift_end"]) if last_out else 0
    early_exit_minutes = max(early_exit, 0)

    working_hours = 0
    if first_in and last_out and last_out > first_in:
        working_hours = round(minutes_between(first_in, last_out) / 60, 2)

    if leave_type and not in_times and not out_times:
        status = "ON_LEAVE"
    elif not first_in or not last_out:
        status = "MISSING_PUNCH"
    elif late_minutes:
        status = "LATE"
    else:
        status = "PRESENT"

    return {
        "employee_code": first["employee_code"],
        "employee_name": first["employee_name"],
        "department": first["department"],
        "department_path": find_department_path(DEPARTMENT_HIERARCHY, first["department"]) or "Unmapped",
        "punch_date": first["punch_date"],
        "first_in": first_in,
        "last_out": last_out,
        "working_hours": working_hours,
        "late_minutes": late_minutes,
        "early_exit_minutes": early_exit_minutes,
        "status": status,
    }


def group_by_employee_day(records):
    groups = {}
    for record in records:
        key = (record["punch_date"], record["employee_code"])
        groups.setdefault(key, []).append(record)
    return dict(sorted(groups.items()))


def summary_anomalies(summary):
    base = {"row_number": "", "employee_code": summary["employee_code"], "punch_date": summary["punch_date"]}
    anomalies = []
    if summary["late_minutes"]:
        anomalies.append({**base, "anomaly_type": "LATE_ARRIVAL", "details": f"Late by {summary['late_minutes']} minutes"})
    if summary["status"] == "MISSING_PUNCH" and not summary["first_in"]:
        anomalies.append({**base, "anomaly_type": "MISSING_IN_PUNCH", "details": "No IN punch recorded"})
    if summary["status"] == "MISSING_PUNCH" and not summary["last_out"]:
        anomalies.append({**base, "anomaly_type": "MISSING_OUT_PUNCH", "details": "No OUT punch recorded"})
    if summary["early_exit_minutes"] and summary["status"] != "MISSING_PUNCH":
        anomalies.append({**base, "anomaly_type": "EARLY_EXIT", "details": f"Left {summary['early_exit_minutes']} minutes early"})
    return anomalies


def correction_row(summary):
    missing = "IN" if not summary["first_in"] else "OUT"
    recorded = summary["last_out"] if missing == "IN" else summary["first_in"]
    return {
        "employee_code": summary["employee_code"],
        "employee_name": summary["employee_name"],
        "department": summary["department"],
        "punch_date": summary["punch_date"],
        "missing_punch": missing,
        "recorded_punch": recorded or "",
        "suggested_action": f"Raise a correction to add the missing {missing} punch",
    }


def write_csv(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


class AttendanceProcessor:
    def __init__(self, path):
        self.path = Path(path)
        self.records = []
        self.anomalies = []
        self.summaries = []
        self.corrections = []
        self.counters = {"total_rows": 0, "valid_rows": 0, "invalid_rows": 0, "duplicate_rows": 0, "leave_rows": 0}
        self.seen_punches = set()

    @log_execution
    def run(self):
        try:
            for row_number, row in stream_records(self.path):
                self.counters["total_rows"] += 1
                self.handle_row(row_number, row)
        except (UnicodeDecodeError, csv.Error, zipfile.BadZipFile, InvalidFileException) as error:
            raise UnsupportedFileError("The file could not be read. Upload a valid CSV or XLSX file.") from error

        if self.counters["total_rows"] == 0:
            raise EmptyFileError("The file has a header row but no attendance rows.")

        self.summaries = [summarize_day(records) for records in group_by_employee_day(self.records).values()]
        for summary in self.summaries:
            self.anomalies.extend(summary_anomalies(summary))
        self.corrections = [correction_row(summary) for summary in self.summaries if summary["status"] == "MISSING_PUNCH"]
        return self

    def handle_row(self, row_number, row):
        try:
            record = validate_row(row_number, row)
        except InvalidRowError as error:
            self.counters["invalid_rows"] += 1
            self.add_row_anomaly(row_number, row, "INVALID_ROW", str(error))
            return

        punch_key = (record["employee_code"], record["punch_date"], record["punch_time"], record["punch_type"])
        if punch_key in self.seen_punches:
            self.counters["duplicate_rows"] += 1
            self.add_row_anomaly(row_number, row, "DUPLICATE_PUNCH", "Same employee, date, time and punch type already seen")
            return

        self.seen_punches.add(punch_key)
        self.records.append(record)
        self.counters["valid_rows"] += 1
        if record["punch_time"] is None:
            self.counters["leave_rows"] += 1

    def add_row_anomaly(self, row_number, row, anomaly_type, details):
        self.anomalies.append({
            "row_number": row_number,
            "employee_code": row.get("employee_code", ""),
            "punch_date": row.get("punch_date", ""),
            "anomaly_type": anomaly_type,
            "details": details,
        })

    def status_counts(self):
        statuses = {summary["status"] for summary in self.summaries}
        return {status: sum(1 for summary in self.summaries if summary["status"] == status) for status in sorted(statuses)}

    @log_execution
    def write_reports(self, output_dir):
        folder = Path(output_dir)
        folder.mkdir(parents=True, exist_ok=True)
        paths = {
            "daily_summary": folder / "daily_summary.csv",
            "anomaly_report": folder / "anomaly_report.csv",
            "correction_required": folder / "correction_required.csv",
        }
        write_csv(paths["daily_summary"], SUMMARY_COLUMNS, self.summaries)
        write_csv(paths["anomaly_report"], ANOMALY_COLUMNS, self.anomalies)
        write_csv(paths["correction_required"], CORRECTION_COLUMNS, self.corrections)
        return paths


def main():
    if len(sys.argv) < 2:
        print("Usage: python attendance/processor.py <attendance_file> [output_folder]")
        return 1

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    output_dir = sys.argv[2] if len(sys.argv) > 2 else "reports/generated"

    try:
        processor = AttendanceProcessor(sys.argv[1]).run()
    except AttendanceFileError as error:
        print(f"Import failed: {error}")
        return 1

    paths = processor.write_reports(output_dir)
    print("Counters:", processor.counters)
    print("Daily status counts:", processor.status_counts())
    for name, path in paths.items():
        print(f"{name}: {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())